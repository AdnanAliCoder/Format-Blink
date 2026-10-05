"""Bounded, single-job conversion worker. Never fetches user-supplied URLs."""
import io
import json
import os
from pathlib import Path
import subprocess
import sys

EXTENSIONS = {
    'pdf-to-word': ({'.pdf'}, 'docx'), 'word-to-pdf': ({'.doc', '.docx', '.odt'}, 'pdf'),
    'ocr-pdf': ({'.pdf'}, 'pdf'), 'pdf-to-excel': ({'.pdf'}, 'xlsx'),
    'excel-to-pdf': ({'.xls', '.xlsx', '.ods'}, 'pdf'), 'pdf-to-powerpoint': ({'.pdf'}, 'pptx'),
    'powerpoint-to-pdf': ({'.ppt', '.pptx', '.odp'}, 'pdf'),
    'protect-pdf': ({'.pdf'}, 'pdf'), 'unlock-pdf': ({'.pdf'}, 'pdf'),
    'html-to-pdf': ({'.html', '.htm'}, 'pdf'),
    'background-remover': ({'.jpg', '.jpeg', '.png', '.webp', '.avif', '.bmp', '.gif'}, 'png'),
    'image-to-text': ({'.jpg', '.jpeg', '.png', '.webp', '.avif', '.bmp', '.gif'}, 'txt'),
    'video-to-text': ({'.mp4', '.webm', '.mov', '.mkv', '.avi', '.m4v'}, 'txt'),
}

def command(args, timeout=180):
    result = subprocess.run(args, capture_output=True, timeout=timeout)
    if result.returncode:
        raise ValueError('The conversion engine could not process this file.')
    return result.stdout

def pdf_reader(source, password=''):
    from pypdf import PdfReader
    reader = PdfReader(source)
    if reader.is_encrypted and not reader.decrypt(password):
        raise ValueError('The PDF password is incorrect or missing.')
    if not 0 < len(reader.pages) <= 100:
        raise ValueError('Use PDFs containing 1–100 pages.')
    return reader

def render(source, directory):
    pdf_reader(source)
    import pypdfium2 as pdfium
    pdf = pdfium.PdfDocument(str(source))
    images = []
    try:
        for index in range(len(pdf)):
            page = pdf[index]
            bitmap = page.render(scale=2.0)
            image = directory / f'page-{index+1}.png'
            bitmap.to_pil().save(image, 'PNG')
            bitmap.close()
            images.append(image)
            page.close()
    finally:
        pdf.close()
    return images

def ocr_language(preferred='eng+urd+hin'):
    import pytesseract
    available = set(pytesseract.get_languages(config=''))
    requested = [name for name in str(preferred or 'eng+urd+hin').split('+') if name in available]
    if not requested:
        requested = [name for name in ('eng', 'urd', 'hin') if name in available]
    if not requested:
        raise ValueError('OCR language data is not installed on the processor.')
    return '+'.join(requested)

def rebuild_ocr_page(image, ocr_pdf, width, height):
    """Create a clean OCR PDF page: editable text plus non-text graphics, no page screenshot."""
    import html
    import fitz
    source = fitz.open(stream=ocr_pdf, filetype='pdf')
    rebuilt = fitz.open()
    try:
        ocr_page = source[0]
        page = rebuilt.new_page(width=width, height=height)
        sx = width / max(1, ocr_page.rect.width)
        sy = height / max(1, ocr_page.rect.height)
        words = ocr_page.get_text('words', sort=False)
        lines = {}
        for word in words:
            if len(word) < 8 or not str(word[4]).strip():
                continue
            lines.setdefault((word[5], word[6]), []).append(word)
        for items in lines.values():
            items.sort(key=lambda item: item[7])
            text = ' '.join(str(item[4]).strip() for item in items if str(item[4]).strip())
            if not text:
                continue
            x0 = min(item[0] for item in items) * sx
            y0 = min(item[1] for item in items) * sy
            x1 = max(item[2] for item in items) * sx
            y1 = max(item[3] for item in items) * sy
            font_size = max(6, (y1-y0) * 0.78)
            rect = fitz.Rect(x0, y0, max(x0+2, x1+4), max(y0+font_size+2, y1+4))
            markup = f'<div style="font-size:{font_size:.2f}pt;line-height:1;margin:0">{html.escape(text)}</div>'
            try:
                page.insert_htmlbox(rect, markup, css='* { font-family: sans-serif; }')
            except Exception:
                latin = text.encode('latin-1', errors='ignore').decode().strip()
                if latin:
                    page.insert_textbox(rect, latin, fontsize=font_size, fontname='helv')

        # Preserve photos, logos and diagrams as separate image crops. OCR text
        # boxes are masked out first so normal paragraphs are not reinserted as pictures.
        try:
            import cv2
            import numpy as np
            pixels = np.array(image.convert('RGB'))
            gray = cv2.cvtColor(pixels, cv2.COLOR_RGB2GRAY)
            mask = cv2.threshold(gray, 242, 255, cv2.THRESH_BINARY_INV)[1]
            ow, oh = max(1, ocr_page.rect.width), max(1, ocr_page.rect.height)
            for word in words:
                x0 = max(0, int(word[0] / ow * image.width) - 5)
                y0 = max(0, int(word[1] / oh * image.height) - 5)
                x1 = min(image.width, int(word[2] / ow * image.width) + 5)
                y1 = min(image.height, int(word[3] / oh * image.height) + 5)
                mask[y0:y1, x0:x1] = 0
            kernel = np.ones((5, 5), np.uint8)
            mask = cv2.morphologyEx(mask, cv2.MORPH_CLOSE, kernel, iterations=2)
            count, _, stats, _ = cv2.connectedComponentsWithStats(mask, 8)
            page_area = image.width * image.height
            minimum = max(1800, int(page_area * 0.0015))
            for index in range(1, count):
                x, y, w, h, area = [int(v) for v in stats[index]]
                box_area = w * h
                if area < minimum or w < 36 or h < 36 or not box_area:
                    continue
                if box_area > page_area * 0.82 or area / box_area < 0.07:
                    continue
                crop = image.crop((x, y, x+w, y+h))
                buffer = io.BytesIO()
                crop.save(buffer, 'PNG')
                rect = fitz.Rect(
                    x / image.width * width,
                    y / image.height * height,
                    (x+w) / image.width * width,
                    (y+h) / image.height * height,
                )
                page.insert_image(rect, stream=buffer.getvalue(), keep_proportion=False)
        except Exception:
            # Graphic detection is best-effort; editable OCR text remains the priority.
            pass
        return rebuilt.tobytes(garbage=4, deflate=True)
    finally:
        source.close()
        rebuilt.close()

def prepare_word_source(source, reader, directory, language):
    import re
    scanned = [
        index for index, page in enumerate(reader.pages)
        if len(re.sub(r'\s+', '', page.extract_text() or '')) < 8
    ]
    if not scanned:
        return source, False
    import pypdfium2 as pdfium
    import pytesseract
    from pypdf import PdfReader, PdfWriter
    lang = ocr_language(language)
    writer = PdfWriter()
    pdf = pdfium.PdfDocument(str(source))
    try:
        for index, original in enumerate(reader.pages):
            if index not in scanned:
                writer.add_page(original)
                continue
            page = pdf[index]
            bitmap = page.render(scale=2.0)
            image = bitmap.to_pil().convert('RGB')
            try:
                data = pytesseract.image_to_pdf_or_hocr(
                    image,
                    extension='pdf',
                    lang=lang,
                    config='--dpi 144 --psm 3',
                    timeout=180,
                )
                rebuilt = rebuild_ocr_page(
                    image,
                    data,
                    float(original.mediabox.width),
                    float(original.mediabox.height),
                )
            finally:
                bitmap.close()
                page.close()
            recognized = PdfReader(io.BytesIO(rebuilt)).pages[0]
            writer.add_page(recognized)
    finally:
        pdf.close()
    target = directory / 'word-ocr-source.pdf'
    with target.open('wb') as handle:
        writer.write(handle)
    return target, True

def strip_full_page_word_images(output):
    # OCR PDFs can contain a full-page scan behind a text layer. Keep the text
    # editable and remove page-sized screenshots while retaining smaller images.
    from docx import Document
    doc = Document(output)
    if not doc.sections:
        return 0
    max_width = max(int(section.page_width) for section in doc.sections)
    max_height = max(int(section.page_height) for section in doc.sections)
    removed = 0
    for shape in list(doc.inline_shapes):
        if int(shape.width) >= int(max_width * 0.82) and int(shape.height) >= int(max_height * 0.55):
            parent = shape._inline.getparent()
            if parent is not None:
                parent.remove(shape._inline)
                removed += 1
    if removed:
        doc.save(output)
    return removed

def ocr_word_fallback(source, output, language):
    # Last-resort scanned-document path: always return editable text rather than
    # silently embedding whole PDF pages as pictures.
    import pypdfium2 as pdfium
    import pytesseract
    from docx import Document
    doc = Document()
    pdf = pdfium.PdfDocument(str(source))
    found = False
    try:
        lang = ocr_language(language)
        for index in range(len(pdf)):
            if index:
                doc.add_page_break()
            page = pdf[index]
            bitmap = page.render(scale=2.0)
            image = bitmap.to_pil().convert('RGB')
            try:
                text = pytesseract.image_to_string(image, lang=lang, config='--psm 3', timeout=180)
            finally:
                bitmap.close()
                page.close()
            chunks = [part.strip() for part in text.split('\n\n') if part.strip()]
            for chunk in chunks:
                doc.add_paragraph(chunk)
                found = True
    finally:
        pdf.close()
    if not found:
        raise ValueError('No readable text could be recognized from this PDF.')
    doc.save(output)

def convert(slug, source, settings, directory):
    if slug not in EXTENSIONS or source.suffix.lower() not in EXTENSIONS[slug][0]:
        raise ValueError('Unsupported input format.')
    output = directory / ('result.' + EXTENSIONS[slug][1])
    language = settings.get('language', 'eng')
    if language not in {'eng', 'urd', 'hin', 'eng+urd', 'eng+hin', 'eng+urd+hin'}:
        raise ValueError('Unsupported OCR language.')
    if slug in {'word-to-pdf', 'excel-to-pdf', 'powerpoint-to-pdf'}:
        profile = directory / 'lo-profile'
        # Dedicated profile disables macros and external link updates.
        user = profile / 'user'
        user.mkdir(parents=True)
        (user / 'registrymodifications.xcu').write_text('''<?xml version="1.0"?><oor:items xmlns:oor="http://openoffice.org/2001/registry"><item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item><item oor:path="/org.openoffice.Office.Calc/Content/Update"><prop oor:name="Link" oor:op="fuse"><value>2</value></prop></item></oor:items>''')
        target = directory / 'office'
        target.mkdir()
        command([__import__('shutil').which('libreoffice') or __import__('shutil').which('soffice') or 'libreoffice', '-env:UserInstallation=' + profile.as_uri(), '--headless', '--convert-to', 'pdf', '--outdir', str(target), str(source)])
        converted = target / (source.stem + '.pdf')
        if not converted.exists():
            raise ValueError('Office conversion failed. Check the document format.')
        converted.replace(output)
    elif slug in {'protect-pdf', 'unlock-pdf'}:
        from pypdf import PdfWriter
        password = settings.get('password', '')
        if not isinstance(password, str) or not 1 <= len(password) <= 128:
            raise ValueError('Enter a password of 1–128 characters.')
        reader = pdf_reader(source, password if slug == 'unlock-pdf' else '')
        writer = PdfWriter(clone_from=reader)
        if slug == 'protect-pdf':
            writer.encrypt(password, algorithm='AES-256')
        with output.open('wb') as handle:
            writer.write(handle)
    elif slug == 'pdf-to-word':
        # Automatic path: searchable PDFs convert directly; scanned/mixed PDFs
        # are OCRed first. The user never has to choose an OCR/appearance mode.
        reader = pdf_reader(source)
        word_source, used_ocr = prepare_word_source(source, reader, directory, language)
        from pdf2docx import Converter
        converter = Converter(str(word_source))
        try:
            converter.convert(str(output), multi_processing=False)
        except Exception:
            if used_ocr:
                ocr_word_fallback(source, output, language)
            else:
                raise ValueError('Editable PDF-to-Word conversion failed for this document.')
        finally:
            converter.close()
        if used_ocr:
            strip_full_page_word_images(output)
        import zipfile
        try:
            with zipfile.ZipFile(output) as package:
                document_xml = package.read('word/document.xml')
        except Exception as error:
            raise ValueError('The Word document could not be validated.') from error
        if b'<w:t' not in document_xml:
            ocr_word_fallback(source, output, language)
    elif slug == 'pdf-to-excel':
        import pdfplumber
        from openpyxl import Workbook
        pdf_reader(source)
        book = Workbook()
        book.remove(book.active)
        with pdfplumber.open(source) as pdf:
            for i, page in enumerate(pdf.pages):
                for j, table in enumerate(page.extract_tables()):
                    sheet = book.create_sheet(f'Page {i+1} table {j+1}')
                    for row in table:
                        sheet.append(row)
                    # User-supplied table text is always text, never spreadsheet formulas.
                    for row in sheet:
                        for cell in row:
                            if isinstance(cell.value, str):
                                cell.data_type = 's'
                    from openpyxl.styles import Alignment
                    from openpyxl.utils import get_column_letter
                    sheet.freeze_panes = 'A2'
                    for row in sheet:
                        for cell in row:
                            cell.alignment = Alignment(wrap_text=True, vertical='top')
                    for column in sheet.columns:
                        length = max((len(str(c.value or '')) for c in column), default=10)
                        sheet.column_dimensions[get_column_letter(column[0].column)].width = min(60, max(12, length+2))
        if not book.sheetnames:
            raise ValueError('No text-based tables detected. OCR scanned pages first.')
        book.save(output)
    elif slug == 'pdf-to-powerpoint':
        # Rebuild text and embedded images as editable slide objects instead of
        # placing a screenshot of each PDF page on the slide.
        import fitz
        import pytesseract
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.util import Pt
        from pytesseract import Output as TesseractOutput
        deck = Presentation()
        pdf = fitz.open(str(source))
        first = pdf[0].rect
        deck.slide_width = Pt(first.width)
        deck.slide_height = Pt(first.height)
        lang = ocr_language(language)
        try:
            for index, page in enumerate(pdf):
                slide = deck.slides.add_slide(deck.slide_layouts[6])
                sx = deck.slide_width / page.rect.width
                sy = deck.slide_height / page.rect.height
                blocks = page.get_text('dict').get('blocks', [])
                has_text = False
                for block in blocks:
                    bbox = block.get('bbox', (0, 0, 0, 0))
                    x0, y0, x1, y1 = bbox
                    if block.get('type') == 0:
                        for line in block.get('lines', []):
                            line_box = line.get('bbox', bbox)
                            lx0, ly0, lx1, ly1 = line_box
                            spans = [span for span in line.get('spans', []) if span.get('text')]
                            if not spans:
                                continue
                            has_text = True
                            shape = slide.shapes.add_textbox(
                                int(lx0 * sx), int(ly0 * sy),
                                max(1, int((lx1-lx0) * sx)), max(1, int((ly1-ly0) * sy))
                            )
                            frame = shape.text_frame
                            frame.clear()
                            frame.margin_left = frame.margin_right = 0
                            frame.margin_top = frame.margin_bottom = 0
                            paragraph = frame.paragraphs[0]
                            for span in spans:
                                run = paragraph.add_run()
                                run.text = span.get('text', '')
                                run.font.size = Pt(max(1, float(span.get('size', 10))))
                                run.font.name = span.get('font') or None
                                flags = int(span.get('flags', 0))
                                run.font.bold = bool(flags & 16)
                                run.font.italic = bool(flags & 2)
                                color = int(span.get('color', 0))
                                run.font.color.rgb = RGBColor((color >> 16) & 255, (color >> 8) & 255, color & 255)
                    elif block.get('type') == 1 and block.get('image'):
                        page_area = max(1, page.rect.width * page.rect.height)
                        block_area = max(0, (x1-x0) * (y1-y0))
                        # Skip full-page scans; OCR below recreates their text.
                        if block_area / page_area < 0.82:
                            slide.shapes.add_picture(
                                io.BytesIO(block['image']),
                                int(x0 * sx), int(y0 * sy),
                                width=max(1, int((x1-x0) * sx)),
                                height=max(1, int((y1-y0) * sy)),
                            )
                if not has_text:
                    pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                    from PIL import Image
                    image = Image.frombytes('RGB', [pix.width, pix.height], pix.samples)
                    data = pytesseract.image_to_data(
                        image, lang=lang, config='--psm 3',
                        output_type=TesseractOutput.DICT, timeout=180
                    )
                    px = page.rect.width / image.width
                    py = page.rect.height / image.height
                    for i, word in enumerate(data.get('text', [])):
                        word = (word or '').strip()
                        try:
                            confidence = float(data['conf'][i])
                        except Exception:
                            confidence = -1
                        if not word or confidence < 20:
                            continue
                        left, top = data['left'][i], data['top'][i]
                        width, height = data['width'][i], data['height'][i]
                        shape = slide.shapes.add_textbox(
                            int(left * px * sx), int(top * py * sy),
                            max(1, int(width * px * sx * 1.15)),
                            max(1, int(height * py * sy * 1.4)),
                        )
                        frame = shape.text_frame
                        frame.clear()
                        frame.margin_left = frame.margin_right = 0
                        frame.margin_top = frame.margin_bottom = 0
                        run = frame.paragraphs[0].add_run()
                        run.text = word
                        run.font.size = Pt(max(6, height * py * 0.8))
        finally:
            pdf.close()
        if not deck.slides:
            raise ValueError('This PDF contains no pages.')
        deck.save(output)
    elif slug in {'ocr-pdf', 'image-to-text'}:
        import pytesseract
        from PIL import Image
        Image.MAX_IMAGE_PIXELS = 25_000_000
        active_language = ocr_language(language)
        if slug == 'image-to-text':
            with Image.open(source) as im:
                if im.width * im.height > 25_000_000:
                    raise ValueError('Use an image under 25 megapixels.')
                text = pytesseract.image_to_string(im, lang=active_language, timeout=120)
            if not text.strip():
                raise ValueError('No text recognized. Try a clearer image or another language.')
            output.write_text(text, encoding='utf-8')
        else:
            import pypdfium2 as pdfium
            from pypdf import PdfWriter, PdfReader
            reader = pdf_reader(source)
            writer = PdfWriter()
            pdf = pdfium.PdfDocument(str(source))
            try:
                for index, original in enumerate(reader.pages):
                    page = pdf[index]
                    bitmap = page.render(scale=2.0)
                    image = bitmap.to_pil().convert('RGB')
                    try:
                        data = pytesseract.image_to_pdf_or_hocr(image, extension='pdf', lang=active_language, config='--dpi 144 --psm 3', timeout=180)
                    finally:
                        bitmap.close()
                        page.close()
                    recognized = PdfReader(io.BytesIO(data)).pages[0]
                    recognized.scale_to(float(original.mediabox.width), float(original.mediabox.height))
                    writer.add_page(recognized)
            finally:
                pdf.close()
            with output.open('wb') as handle:
                writer.write(handle)
    elif slug == 'html-to-pdf':
        html = source.read_text(encoding='utf-8')
        if any(token in html.lower() for token in ('<script', 'http://', 'https://', 'file://')):
            raise ValueError('Use self-contained HTML without scripts or external resources.')
        from weasyprint import HTML
        def deny_resources(url, *args, **kwargs):
            # Permit only embedded raster images; never fetch network or local files.
            from weasyprint import default_url_fetcher
            if url.startswith(('data:image/png;', 'data:image/jpeg;', 'data:image/gif;', 'data:image/webp;')) and len(url) <= 14_000_000:
                return default_url_fetcher(url)
            raise ValueError('External and local resources are disabled.')
        try:
            HTML(string=html, url_fetcher=deny_resources).write_pdf(output)
        except Exception as error:
            raise ValueError('HTML layout rendering failed. Check the HTML and the WeasyPrint installation.') from error
    elif slug == 'background-remover':
        from PIL import Image
        from rembg import remove, new_session
        Image.MAX_IMAGE_PIXELS = 25_000_000
        with Image.open(source) as image:
            if image.width * image.height > 25_000_000:
                raise ValueError('Use an image under 25 megapixels.')
            result = remove(image.convert('RGBA'), session=new_session('u2netp'))
            result.save(output, 'PNG')
    elif slug == 'video-to-text':
        from faster_whisper import WhisperModel
        info = json.loads(command(['ffprobe', '-v', 'error', '-show_format', '-show_streams', '-of', 'json', str(source)]))
        if not any(s.get('codec_type') == 'audio' for s in info.get('streams', [])):
            raise ValueError('This video has no audio track.')
        if float(info['format']['duration']) > 600:
            raise ValueError('Transcribe clips up to 10 minutes here. Use Clip Studio for longer recordings.')
        model = WhisperModel(os.environ.get('WHISPER_MODEL', 'small'), device='cpu', compute_type='int8', cpu_threads=2)
        # Decode with the installed FFmpeg instead of version-sensitive PyAV APIs.
        import numpy as np
        pcm = command(['ffmpeg', '-v', 'error', '-i', str(source), '-vn', '-ac', '1', '-ar', '16000', '-f', 'f32le', 'pipe:1'])
        audio = np.frombuffer(pcm, dtype=np.float32).copy()
        segments, _ = model.transcribe(audio, beam_size=3, vad_filter=True)
        text = '\n'.join(f'[{s.start:.2f} – {s.end:.2f}] {s.text.strip()}' for s in segments)
        if not text.strip():
            raise ValueError('No speech detected in this video.')
        output.write_text(text, encoding='utf-8')
    if not output.exists() or output.stat().st_size == 0:
        raise ValueError('No output was produced.')
    return output

if __name__ == '__main__':
    try:
        slug, source = sys.argv[1:3]
        print(convert(slug, Path(source), json.load(sys.stdin), Path(source).parent))
    except Exception as error:
        # Only a concise error reaches the API. No file contents or passwords are logged.
        print(str(error)[:300], file=sys.stderr)
        sys.exit(1)
