import sys,tempfile,subprocess
from pathlib import Path
sys.path.insert(0,str(Path.cwd()/'processor/tools'))
from convert import convert
from reportlab.pdfgen.canvas import Canvas
from PIL import Image
from docx import Document
from pypdf import PdfReader
root=Path('/tmp/formatblink-fidelity');root.mkdir(exist_ok=True)
im=Image.new('RGB',(180,90),'#128ab3');im.save(root/'image.png')
c=Canvas(str(root/'input.pdf'),pagesize=(400,500))
for i in range(2):
 c.setFont('Helvetica-Bold',22);c.drawString(30,450,'Page '+str(i+1));c.drawImage(str(root/'image.png'),30,280,180,90);c.setFont('Helvetica',12);c.drawString(30,230,'Text with image and original layout');c.showPage()
c.save()
d=root/'editable';d.mkdir(exist_ok=True)
out=convert('pdf-to-word',root/'input.pdf',{},d)
doc=Document(out)
text='\n'.join([p.text for p in doc.paragraphs]+[cell.text for table in doc.tables for row in table.rows for cell in row.cells])
assert 'Text with image and original layout' in text
assert len(doc.inline_shapes)>=1
subprocess.run(['soffice','--headless','--convert-to','pdf','--outdir',str(d),str(out)],check=True,capture_output=True)
pages=PdfReader(d/'result.pdf').pages
print('editable','pages',len(pages),'images',len(doc.inline_shapes))
assert len(pages)==2
assert abs(float(pages[0].mediabox.width)-400)<1
subprocess.run(['pdftoppm','-f','1','-singlefile','-scale-to','700','-png',str(d/'result.pdf'),str(d/'preview')],check=True)
