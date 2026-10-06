import sys, tempfile, json, io, os
from pathlib import Path
sys.path.insert(0,str(Path.cwd()/'processor/tools'))
from convert import convert
from reportlab.pdfgen.canvas import Canvas
from reportlab.lib.utils import ImageReader
from pypdf import PdfReader
from PIL import Image,ImageDraw,ImageFont
from docx import Document
from openpyxl import Workbook,load_workbook
from pptx import Presentation
root=Path(os.environ.get('FIXTURE_DIR','/tmp/formatblink-fixtures'));root.mkdir(exist_ok=True)
c=Canvas(str(root/'sample.pdf'),pagesize=(400,500));c.drawString(40,450,'FormatBlink Test SECRET 123');
for x in [40,150,260]: c.line(x,250,x,350)
for y in [250,300,350]: c.line(40,y,260,y)
c.drawString(50,320,'Name');c.drawString(160,320,'Score');c.drawString(50,270,'Adnan');c.drawString(160,270,'93');c.showPage();c.drawString(40,450,'Second page');c.save()
im=Image.new('RGB',(1000,300),'white');d=ImageDraw.Draw(im);d.text((30,70),'FORMAT BLINK TEST 123',font=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',48),fill='black');d.rectangle((700,50,930,250),fill=(20,130,190));im.save(root/'text.png')
c=Canvas(str(root/'scanned.pdf'),pagesize=(1000,300));c.drawImage(str(root/'text.png'),0,0,width=1000,height=300);c.save()
doc=Document();doc.add_paragraph('FormatBlink Word fixture');doc.save(root/'sample.docx')
w=Workbook();w.active.append(['Name','Score']);w.active.append(['Adnan',93]);w.save(root/'sample.xlsx')
p=Presentation();p.slides.add_slide(p.slide_layouts[5]).shapes.title.text='FormatBlink presentation';p.save(root/'sample.pptx')
(root/'sample.html').write_text('<h1>FormatBlink HTML</h1><p>Conversion verified</p>')
results={}
for slug,name in [('pdf-to-word','sample.pdf'),('pdf-to-excel','sample.pdf'),('pdf-to-powerpoint','sample.pdf'),('word-to-pdf','sample.docx'),('excel-to-pdf','sample.xlsx'),('powerpoint-to-pdf','sample.pptx'),('protect-pdf','sample.pdf'),('ocr-pdf','scanned.pdf'),('image-to-text','text.png'),('html-to-pdf','sample.html')]:
 try:
  with tempfile.TemporaryDirectory() as tmp:
   out=convert(slug,root/name,{'password':'sample-test-123','language':'eng'},Path(tmp))
   assert out.stat().st_size>0
   if slug=='pdf-to-word': assert 'SECRET 123' in '\n'.join(p.text for p in Document(out).paragraphs)
   if slug=='pdf-to-excel': assert load_workbook(out).active['A2'].value=='Adnan'
   if slug=='pdf-to-powerpoint': assert len(Presentation(out).slides)==2
   if slug=='protect-pdf':
    r=PdfReader(out);assert r.is_encrypted and r.decrypt('sample-test-123');locked=root/'locked.pdf';locked.write_bytes(out.read_bytes())
   elif out.suffix=='.pdf':
    r=PdfReader(out);assert len(r.pages)>0
    if slug=='ocr-pdf': assert 'FORMAT' in r.pages[0].extract_text()
   if slug=='image-to-text': assert 'FORMAT' in out.read_text()
   results[slug]='PASS'
 except Exception as e: results[slug]='FAIL '+str(e)
with tempfile.TemporaryDirectory() as tmp:
 out=convert('unlock-pdf',root/'locked.pdf',{'password':'sample-test-123'},Path(tmp));assert not PdfReader(out).is_encrypted;results['unlock-pdf']='PASS'
 try: convert('unlock-pdf',root/'locked.pdf',{'password':'wrong'},Path(tmp));raise AssertionError('wrong password accepted')
 except ValueError: results['wrong-password']='PASS'
with tempfile.TemporaryDirectory() as tmp:
 out=convert('pdf-to-word',root/'scanned.pdf',{'language':'eng+urd+hin'},Path(tmp))
 scanned_doc=Document(out)
 scanned_text='\n'.join(p.text for p in scanned_doc.paragraphs)
 assert 'FORMAT' in scanned_text and '123' in scanned_text
 assert len(scanned_doc.inline_shapes)>=1
 results['pdf-to-word-scanned-auto-ocr']='PASS'
with tempfile.TemporaryDirectory() as tmp:
 out=convert('pdf-to-excel',root/'scanned.pdf',{'language':'eng+urd+hin'},Path(tmp))
 scanned_book=load_workbook(out)
 assert scanned_book.sheetnames
 scanned_values=' '.join(str(cell.value or '') for sheet in scanned_book.worksheets for row in sheet.iter_rows() for cell in row)
 assert 'FORMAT' in scanned_values and '123' in scanned_values
 results['pdf-to-excel-scanned-auto-ocr']='PASS'
with tempfile.TemporaryDirectory() as tmp:
 out=convert('pdf-to-powerpoint',root/'sample.pdf',{'language':'eng'},Path(tmp))
 deck=Presentation(out)
 slide_text='\n'.join(shape.text for slide in deck.slides for shape in slide.shapes if getattr(shape,'has_text_frame',False))
 assert 'FormatBlink Test' in slide_text and 'Second page' in slide_text
 results['pdf-to-powerpoint-editable-text']='PASS'
print(json.dumps(results,indent=2));(root/'processor-results.json').write_text(json.dumps(results,indent=2))

if any(v.startswith('FAIL') for v in results.values()): raise SystemExit(1)
