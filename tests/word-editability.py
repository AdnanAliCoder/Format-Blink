"""Regression fixtures for OCR-layer detection and editable Word output cleanup."""
import importlib.util
import io
from pathlib import Path
import tempfile
import unittest
import fitz
from PIL import Image
from docx import Document
from docx.shared import Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

spec = importlib.util.spec_from_file_location('worker', Path(__file__).parents[1] / 'processor/tools/convert.py')
worker = importlib.util.module_from_spec(spec)
spec.loader.exec_module(worker)

class WordEditability(unittest.TestCase):
    def test_normal_text_and_corner_logo_do_not_force_ocr(self):
        with fitz.open() as pdf:
            page = pdf.new_page()
            page.insert_text((50, 80), 'Editable native text stays native')
            data = io.BytesIO(); Image.new('RGB', (40, 40), 'blue').save(data, 'PNG')
            page.insert_image(fitz.Rect(20, 20, 40, 40), stream=data.getvalue())
            self.assertFalse(worker.page_needs_word_ocr(page))

    def test_full_page_scan_with_existing_text_is_rebuilt(self):
        with fitz.open() as pdf:
            page = pdf.new_page()
            data = io.BytesIO(); Image.new('RGB', (595, 842), 'white').save(data, 'PNG')
            page.insert_image(page.rect, stream=data.getvalue())
            page.insert_text((50, 80), 'Existing OCR layer must not duplicate the scan')
            self.assertTrue(worker.page_needs_word_ocr(page))

    def test_hidden_ocr_text_is_rebuilt(self):
        with fitz.open() as pdf:
            page = pdf.new_page()
            page.insert_text((50, 80), 'Invisible text layer', render_mode=3)
            self.assertTrue(worker.page_needs_word_ocr(page))

    def test_inline_and_floating_screenshots_removed_logo_and_text_kept(self):
        with tempfile.TemporaryDirectory() as directory:
            directory = Path(directory)
            pic = directory / 'fixture.png'; Image.new('RGB', (100, 100), 'blue').save(pic)
            output = directory / 'result.docx'
            doc = Document(); doc.add_paragraph('This text must remain editable')
            doc.add_picture(str(pic), width=Inches(8), height=Inches(10))
            shape = doc.add_picture(str(pic), width=Inches(8), height=Inches(10))
            shape._inline.tag = qn('wp:anchor')
            doc.add_picture(str(pic), width=Inches(.5), height=Inches(.5))
            doc.settings.element.append(OxmlElement('w:documentProtection'))
            doc.settings.element.append(OxmlElement('w:writeProtection'))
            doc.save(output)
            self.assertEqual(worker.strip_full_page_word_images(output), 2)
            clean = Document(output)
            self.assertIn('This text must remain editable', '\n'.join(p.text for p in clean.paragraphs))
            self.assertEqual(len(clean.inline_shapes), 1)
            self.assertEqual(len(list(clean.settings.element.iter(qn('w:documentProtection')))), 0)
            self.assertEqual(len(list(clean.settings.element.iter(qn('w:writeProtection')))), 0)

if __name__ == '__main__': unittest.main()
