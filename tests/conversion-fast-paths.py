"""Native/sparse Word pages and borderless/text Excel regressions."""
import sys
import tempfile
from pathlib import Path
sys.path.insert(0, str(Path(__file__).parents[1]/'processor/tools'))
from convert import convert, page_needs_word_ocr
import fitz
from docx import Document
from openpyxl import load_workbook
from reportlab.pdfgen.canvas import Canvas

with tempfile.TemporaryDirectory() as tmp:
    root = Path(tmp)
    source = root/'borderless.pdf'
    canvas = Canvas(str(source))
    for y, left, right in [(750,'Name','Score'),(720,'Adnan','93'),(690,'Amir','84')]:
        canvas.drawString(50,y,left)
        canvas.drawString(250,y,right)
    canvas.showPage()
    canvas.drawString(50,750,'This readable page must not disappear.')
    canvas.showPage()
    canvas.drawString(50,750,'Hi')
    canvas.save()
    with fitz.open(source) as pdf:
        assert not page_needs_word_ocr(pdf[2]), 'Short native text triggered unnecessary OCR'
    excel = root/'excel';excel.mkdir()
    out = convert('pdf-to-excel', source, {}, excel)
    book = load_workbook(out)
    values = [list(sheet.values) for sheet in book.worksheets]
    assert any(('Adnan','93') in rows for rows in values), values
    assert any('This readable page must not disappear.' in str(rows) for rows in values), values
    assert any('Hi' in str(rows) for rows in values), values
    word = root/'word';word.mkdir()
    out = convert('pdf-to-word', source, {}, word)
    doc = Document(out)
    text = '\n'.join(node.text or '' for node in doc.element.iter() if node.tag.endswith('}t'))
    assert 'Hi' in text and 'Adnan' in text, text
print('PASS: borderless Excel columns, readable page retention, sparse editable Word without OCR')
